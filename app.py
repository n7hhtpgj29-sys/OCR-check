import io,os,json,hashlib
from pathlib import Path
import fitz
import streamlit as st
from openpyxl import Workbook
from engine import extract,load_profiles,FIELDS,NUMERIC

BASE=Path(os.environ.get('CUSTOMS_OCR_BASE',Path(__file__).resolve().parent))
PROFILES=BASE/'templates'
PROFILES.mkdir(exist_ok=True)
st.set_page_config(page_title='Customs OCR v2.0',layout='wide',page_icon='📄')
st.title('Customs OCR v2.0 · Multi-Forwarder')
st.caption('Local PDF extraction · Forwarder/Template matching · Last-page summary · Manual approval · Excel')

def export_excel(rows,logs):
    wb=Workbook();ws=wb.active;ws.title='Upload_Data'
    cols=['Source_File','Forwarder','Template_ID','Pages']+FIELDS+['Review_Status']
    ws.append(cols)
    for row in rows:ws.append([row.get(k,'') for k in cols])
    log=wb.create_sheet('Validation_Log');log.append(['Source_File','Field','Issue','Page'])
    for x in logs:log.append([x.get(k,'') for k in ['Source_File','Field','Issue','Page']])
    from openpyxl.styles import Font,PatternFill
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174A70')
        for col in sheet.columns:
            letter=col[0].column_letter;sheet.column_dimensions[letter].width=min(55,max(18,max(len(str(c.value or '')) for c in list(col)[:100])+2))
    out=io.BytesIO();wb.save(out);return out.getvalue()

profiles=load_profiles(PROFILES)
with st.sidebar:
    st.header('Template Management')
    st.write(f'Active templates: {len(profiles)}')
    for p in profiles:st.caption(f"{p.get('forwarder')} / {p.get('template_id')} / {p.get('version','1')}")
    new=st.file_uploader('Import template JSON',type=['json'],key='profile')
    if new and st.button('Save template'):
        try:
            obj=json.loads(new.getvalue());tid=obj['template_id']
            if not tid.replace('-','').replace('_','').isalnum():raise ValueError('Invalid template ID')
            if not obj.get('forwarder_aliases'):raise ValueError('Forwarder aliases required')
            dest=PROFILES/(tid+'.json')
            if dest.exists():st.warning('Existing template: rename template_id to create a new version')
            else:dest.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8');st.success('Saved; refresh page to activate')
        except Exception as e:st.error(str(e))
    st.download_button('Download example profile',json.dumps({'template_id':'FORWARDER-A-v1','version':'1','forwarder':'Forwarder A','forwarder_aliases':['FORWARDER A CO., LTD.'],'layout_anchors':['เลขที่ใบขนฯ'],'active':True,'rules':{'Invoice_No':{'labels':['Invoice No.'],'pattern':r'([A-Z0-9/-]+)','scope':'first'}}},indent=2,ensure_ascii=False),'example_profile.json','application/json')

files=st.file_uploader('Upload Import Declaration PDFs',type=['pdf'],accept_multiple_files=True)
if files:
    signature=tuple((f.name,hashlib.sha256(f.getvalue()).hexdigest()) for f in files)
    if st.session_state.get('signature')!=signature:
        st.session_state.signature=signature;st.session_state.results=[];st.session_state.approved=False
        for f in files:
            try:
                row,proof,issues=extract(f.getvalue(),f.name,profiles)
                st.session_state.results.append({'row':row,'proof':proof,'issues':issues,'data':f.getvalue()})
            except Exception as e:st.session_state.results.append({'error':str(e),'name':f.name})
    st.metric('PDFs processed',sum('row' in x for x in st.session_state.results))
    for n,item in enumerate(st.session_state.results):
        if 'error' in item:st.error(f"{item['name']}: {item['error']}");continue
        row=item['row']
        with st.expander(f"{row['Source_File']} · {row['Forwarder']} · {row['Template_ID']} · {row['Pages']} pages",expanded=True):
            if row['Forwarder']=='Unknown':st.warning('Unknown forwarder or ambiguous layout: manual review required')
            for field in FIELDS:
                opts=item['proof'].get(field,[])
                hint='; '.join(f"p{x['page']}: {x['value'][:70]}" for x in opts[:3]) or 'No matching field label'
                row[field]=st.text_input(field,value=str(row.get(field,'')),help='Candidates: '+hint,key=f'{n}_{field}',on_change=lambda:st.session_state.update(approved=False))
                if opts:st.caption('Source candidate(s): '+hint)
            if item['issues']:st.warning('Validation warnings: '+str(len(item['issues'])))
            with st.popover('Preview PDF pages'):
                pdf=fitz.open(stream=item['data'],filetype='pdf')
                pn=st.number_input('Page',1,len(pdf),1,key=f'page{n}')
                pix=pdf[int(pn)-1].get_pixmap(matrix=fitz.Matrix(1.25,1.25))
                st.image(pix.tobytes('png'))
    issues=[{'Source_File':x['row']['Source_File'],**issue} for x in st.session_state.results if 'row' in x for issue in x['issues']]
    if issues:
        st.subheader('Validation Log')
        st.table(issues)
    st.warning('Check each field against its PDF. For CIF, Duty and KGM confirm the declaration-level TOTAL on the last page; never sum page subtotals automatically.')
    if st.button('Approve reviewed documents'):
        rows=[x['row'] for x in st.session_state.results if 'row' in x]
        missing=[f"{r['Source_File']}: {key}" for r in rows for key in FIELDS if not str(r.get(key,'')).strip()]
        dup=[r['Declaration_No'] for r in rows]
        if not rows:st.error('No valid documents')
        elif missing:st.error('Missing fields: '+', '.join(missing[:20]))
        elif len(dup)!=len(set(dup)):st.error('Duplicate Declaration No.')
        else:st.session_state.approved=True
    if st.session_state.get('approved'):
        rows=[]
        for item in st.session_state.results:
            if 'row' not in item:continue
            row=dict(item['row']);row['Review_Status']='Approved'
            for key in NUMERIC:
                try:row[key]=float(str(row[key]).replace(',',''))
                except ValueError:pass
            rows.append(row)
        st.download_button('Download Excel',export_excel(rows,issues),'Customs_Import_Export.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
else:st.info('Upload one or more PDF files to start.')
